#!/usr/bin/env python3
"""Build a single-sheet Excel workbook from an enriched Fortify findings JSON.

Usage:
    python3 build_xlsx.py <enriched.json> [-o report.xlsx]

Input JSON schema (produced by extract.py, then enriched by Claude with a
"decision" and "reason" on each finding):

    {
      "meta": { project, commit, scan_date, sca_version, filter_set, machine,
                files_scanned, loc, total_issues, counts{...}, source_pdf },
      "findings": [ { no, level, category, location, owasp, kingdom, abstract,
                      source, sink, decision, reason }, ... ]
    }

Layout: one worksheet with a metadata header block on top (the report
attributes that define what the scan covers) followed by the findings table.
"""
import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LEVEL_FILL = {
    "Critical": "C00000",  # dark red
    "High": "ED7D31",      # orange
    "Medium": "FFC000",    # amber
    "Low": "70AD47",       # green
}
LEVEL_ORDER = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
META_LABEL_FONT = Font(bold=True, color="1F4E78")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")

COLUMNS = [
    ("No.", "no", 6),
    ("Risk Level", "level", 12),
    ("Category", "category", 34),
    ("Location", "location", 40),
    ("OWASP 2021", "owasp", 26),
    ("Abstract", "abstract", 60),
    ("Suggested Decision", "decision", 18),
    ("Reason (draft)", "reason", 60),
    ("Status / Notes", "_notes", 24),
]


def write_meta(ws, meta):
    row = 1
    ws.cell(row, 1, f"Fortify Security Report — {meta.get('project') or ''}").font = TITLE_FONT
    row += 2

    counts = meta.get("counts", {})
    fields = [
        ("Project", meta.get("project")),
        ("Commit", meta.get("commit")),
        ("Scan Date", meta.get("scan_date")),
        ("Fortify SCA Version", meta.get("sca_version")),
        ("Filter Set", meta.get("filter_set")),
        ("Scan Machine", meta.get("machine")),
        ("Files Scanned", meta.get("files_scanned")),
        ("Lines of Code", f"{meta.get('loc'):,}" if meta.get("loc") else None),
        ("Total Issues", meta.get("total_issues")),
        ("Critical / High / Medium / Low",
         f"{counts.get('Critical', 0)} / {counts.get('High', 0)} / "
         f"{counts.get('Medium', 0)} / {counts.get('Low', 0)}"),
        ("Source PDF", meta.get("source_pdf")),
    ]
    for label, value in fields:
        ws.cell(row, 1, label).font = META_LABEL_FONT
        ws.cell(row, 1).alignment = TOP
        c = ws.cell(row, 2, "" if value is None else value)
        c.alignment = TOP
        row += 1
    return row + 1  # leave a blank spacer row


def write_findings(ws, findings, start_row):
    # header
    for col, (title, _key, _w) in enumerate(COLUMNS, start=1):
        c = ws.cell(start_row, col, title)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        c.border = BORDER

    ordered = sorted(
        findings,
        key=lambda f: (LEVEL_ORDER.get(f.get("level"), 9), f.get("no", 0)),
    )

    r = start_row + 1
    for f in ordered:
        for col, (_title, key, _w) in enumerate(COLUMNS, start=1):
            value = "" if key == "_notes" else f.get(key, "")
            c = ws.cell(r, col, value)
            c.alignment = TOP_WRAP
            c.border = BORDER
            if key == "level":
                fill = LEVEL_FILL.get(f.get("level"))
                if fill:
                    c.fill = PatternFill("solid", fgColor=fill)
                    c.font = WHITE_BOLD
                    c.alignment = Alignment(vertical="top", horizontal="center")
        r += 1
    return start_row, r - 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("json")
    ap.add_argument("-o", "--out", default=None)
    args = ap.parse_args()

    with open(args.json) as fh:
        data = json.load(fh)
    meta = data.get("meta", {})
    findings = data.get("findings", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Fortify Report"

    for col, (_t, _k, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    table_start = write_meta(ws, meta)
    header_row, last_row = write_findings(ws, findings, table_start)

    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(COLUMNS))}{last_row}"
    )

    out = args.out or (os.path.splitext(args.json)[0].replace(".findings", "") + ".xlsx")
    wb.save(out)
    print(f"Wrote {out} ({len(findings)} findings).")


if __name__ == "__main__":
    main()
